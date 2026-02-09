"""
📊 Creating Charts in Excel using openpyxl (LineChart)

This script demonstrates how to:
- Load an existing Excel file
- Select a specific range of cells as chart data
- Create a Line Chart using openpyxl
- Customize chart properties (title, axes, style)
- Insert the chart into the Excel worksheet

🎯 Learning Goals
- Understand how Reference objects define chart data ranges
- Learn how to create and configure a LineChart
- Practice adding charts to an Excel worksheet programmatically
"""

# ==================================================
# 1️⃣ Import required libraries
# ==================================================
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference


# ==================================================
# 2️⃣ Load existing Excel file
# ==================================================
# Load sample.xlsx as a Workbook object
wb = load_workbook("sample.xlsx")

# Select the active worksheet
ws = wb.active


# ==================================================
# 3️⃣ Define data range for the chart
# ==================================================
"""
📌 Expected data structure
Column A : ID / Number
Column B : English score
Column C : Math score

- Row 1 contains headers
- Rows 2–11 contain actual data
"""

# Use B1:C11 as the chart data range (including headers)
line_value = Reference(
    ws,
    min_row=1,   # Include header row
    max_row=11,
    min_col=2,   # Column B
    max_col=3    # Column C
)


# ==================================================
# 4️⃣ Create LineChart and add data
# ==================================================
# Create a Line Chart object
line_chart = LineChart()

# Add data to the chart
# titles_from_data=True uses the first row as series names
line_chart.add_data(line_value, titles_from_data=True)


# ==================================================
# 5️⃣ Configure chart settings
# ==================================================
# Set chart title
line_chart.title = "Score Report"

# Apply a predefined chart style (1–48)
line_chart.style = 10

# Set axis titles
line_chart.y_axis.title = "Score"
line_chart.x_axis.title = "ID"


# ==================================================
# 6️⃣ Insert chart into worksheet
# ==================================================
# Insert the chart at position E1
ws.add_chart(line_chart, "E1")


# ==================================================
# 7️⃣ Save the Excel file
# ==================================================
# Save as a new file with the chart included
wb.save("sample_chart.xlsx")

print("✅ Excel file with LineChart has been created successfully.")
